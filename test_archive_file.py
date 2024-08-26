from zipfile import ZipFile
import csv

from pypdf import PdfReader
from openpyxl import load_workbook

from script_os import ARCHIVE_FILE

files = ['test_task.xlsx', 'test_task.csv', 'test_task.pdf']

def normalize_text(text):
    return ' '.join(text.split())


def test_archive_csv():
    with ZipFile(ARCHIVE_FILE, 'r') as zip_file:
        with zip_file.open('test_task.csv') as csv_file:
            csv_content = csv_file.read().decode('utf-8-sig')
            csvreader = list(csv.reader(csv_content.splitlines()))
            row_1 = csvreader[1]

            assert row_1[0] == 'Условия '


def test_xlsx():
    with ZipFile(ARCHIVE_FILE, 'r') as zip_file:
        with zip_file.open('test_task.xlsx') as xlsx_file:
            workbook = load_workbook(xlsx_file)
            sheet = workbook.active
            cell_text = sheet.cell(row=3, column=1).value

            assert "Употребление алкоголя <= 17 единиц в неделю" in cell_text


def test_pdf():
    with ZipFile(ARCHIVE_FILE, 'r') as zip_file:
        with zip_file.open('test_task.pdf') as pdf_file:
            reader = PdfReader(pdf_file)
            page_text = reader.pages[1].extract_text()

            expected_text = (
                "Тест-кейс 30ИМТ =  30,01 не курит, "
                "заполнение анкеты "
                '"Оценка риска для '
                'здоровья"Сотрудник должен '
                'получить $69 скидку '
                'на оплату '
                'Тест-кейс 31ИМТ = 25,5 курит участвует '
                'в курсе отказа от курения, '
                'заполнение анкеты '
                '"Оценка риска для '
                'здоровья"Сотрудник должен '
                'получить $97 скидку '
                'на оплату '
                'Тест-кейс 32ИМТ =  29,98 курит '
                'участвует в курсе отказа от '
                'курения, заполнение '
                'анкеты "Оценка риска для '
                'здоровья"Сотрудник должен '
                'получить $66 скидку '
                'на оплату '
                'Тест-кейс 33ИМТ = 30,02 курит '
                'участвует в курсе отказа от '
                'курения, заполнение '
                'анкеты "Оценка риска для '
                'здоровья"Сотрудник должен '
                'получить $47 скидку '
                'на оплату '
                'Тест-кейс 34ИМТ = 25,49 курит не '
                'участвует в курсе отказа от '
                'курения, заполнение '
                'анкеты "Оценка риска для '
                'здоровья"Сотрудник должен '
                'доплатить еще $2 '
                'Тест-кейс 35ИМТ =  25,51 курит не '
                'участвует в курсе отказа от '
                'курения, заполнение '
                'анкеты "Оценка риска для '
                'здоровья"Сотрудник должен '
                'доплатить еще $33 '
                'Тест-кейс 36ИМТ = 30 курит не '
                'участвует в курсе отказа от '
                'курения, заполнение '
                'анкеты "Оценка риска для '
                'здоровья"Сотрудник должен '
                'доплатить еще $52'
            )

            normalized_page_text = normalize_text(page_text)
            normalized_expected_text = normalize_text(expected_text)

            assert normalized_expected_text in normalized_page_text, "Текст не найден в PDF!"
